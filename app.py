import re
from io import BytesIO

import pandas as pd
import streamlit as st

# =========================================================
# PAGE CONFIG
# =========================================================
st.set_page_config(
    page_title="Smart Excel Reporting Agent",
    page_icon="ðŸ“Š",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =========================================================
# PROFESSIONAL UI
# =========================================================
st.markdown(
    """
    <style>
    .stApp { background: #f6f8fb; }
    [data-testid="stHeader"] { background: transparent; }
    .main-title {
        font-size: 38px;
        font-weight: 800;
        color: #12355b;
        margin: 0 0 2px 0;
    }
    .sub-title {
        color: #64748b;
        font-size: 16px;
        margin-bottom: 18px;
    }
    .section-title {
        color: #12355b;
        font-weight: 750;
    }
    .info-card {
        background: white;
        border: 1px solid #dbe3ee;
        border-radius: 12px;
        padding: 14px 16px;
        margin-bottom: 10px;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.04);
    }
    div[data-testid="stMetric"] {
        background: white;
        border: 1px solid #dbe3ee;
        border-radius: 12px;
        padding: 12px 14px;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.04);
    }
    .small-note { color: #64748b; font-size: 13px; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown('<div class="main-title">ðŸ“Š Smart Excel Reporting Agent</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="sub-title">Professional multi-file Excel consolidation, validation and reporting system</div>',
    unsafe_allow_html=True,
)

# =========================================================
# CONSTANTS / HELPERS
# =========================================================
SOURCE_COLUMNS = ["Source File", "Source Sheet"]
MISSING_TEXT = {"", "none", "nan", "null", "na", "n/a", "-"}


def norm_key(value):
    text = str(value).replace("\n", " ").replace("\r", " ").strip().casefold()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text


def normalize_columns(columns):
    mapping = {
        "uin": "UIN", "employee_number": "UIN", "employee_no": "UIN",
        "employee_id": "UIN", "employee_code": "UIN",
        "name": "Employee Name", "employee": "Employee Name", "employee_name": "Employee Name",
        "bn": "BN", "battalion": "BN", "battalion_no": "BN", "battalion_number": "BN",
        "unit": "BN", "unit_name": "BN",
        "pran": "PRAN No", "pran_no": "PRAN No", "pran_number": "PRAN No",
        "status": "Case Status", "case_status": "Case Status",
        "amount": "Contribution", "contribution": "Contribution", "total_amount": "Contribution",
        "nomination": "Nomination", "nomination_status": "Nomination",
        "nominee": "Nominee", "nominee_status": "Nominee",
        "mobile": "Mobile No", "mobile_no": "Mobile No", "mobile_number": "Mobile No",
        "phone": "Mobile No", "phone_no": "Mobile No",
        "date": "Date", "transaction_date": "Date", "payment_date": "Date",
    }
    result = []
    counts = {}
    for col in columns:
        original = str(col).strip()
        base = mapping.get(norm_key(original), original or "Unnamed Column")
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def merge_duplicate_columns(df):
    groups = {}
    for i, col in enumerate(df.columns):
        base = re.sub(r"_\d+$", "", str(col))
        groups.setdefault(base, []).append(i)
    drop = []
    for base, positions in groups.items():
        if len(positions) < 2:
            continue
        combined = pd.Series(pd.NA, index=df.index, dtype="object")
        for pos in positions:
            s = df.iloc[:, pos].replace({"": pd.NA, "None": pd.NA, "none": pd.NA, "nan": pd.NA, "NaN": pd.NA})
            combined = combined.fillna(s)
        df.iloc[:, positions[0]] = combined
        drop.extend(positions[1:])
    if drop:
        df = df.iloc[:, [i for i in range(len(df.columns)) if i not in set(drop)]].copy()
    return df


def clean_blank_rows(df):
    if df.empty:
        return df.copy()
    check = df.copy()
    check = check.replace(r"^\s*$", pd.NA, regex=True)
    check = check.replace({"None": pd.NA, "none": pd.NA, "nan": pd.NA, "NaN": pd.NA, "null": pd.NA, "NULL": pd.NA})
    return df.loc[~check.isna().all(axis=1)].reset_index(drop=True)


def canonical_value(value, column_name=""):
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    low = text.casefold()
    if low in MISSING_TEXT:
        return ""
    if "date" in norm_key(column_name):
        parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
        if not pd.isna(parsed):
            return parsed.strftime("%Y-%m-%d")
    if "contribution" in norm_key(column_name) or "amount" in norm_key(column_name):
        num = pd.to_numeric(text.replace(",", ""), errors="coerce") if isinstance(text, str) else pd.to_numeric(text, errors="coerce")
        if not pd.isna(num):
            return f"{float(num):.2f}"
    return re.sub(r"\s+", " ", text).casefold()


def deduplicate_exact_business_records(df):
    """Remove the same business record repeated in multiple files/sheets.
    Source File/Source Sheet are ignored. Values are normalized so that
    Excel date/number/text representation differences do not create copies.
    """
    if df.empty:
        return df.copy()
    compare_cols = [c for c in df.columns if c not in SOURCE_COLUMNS]
    if not compare_cols:
        return df.reset_index(drop=True)
    keys = pd.DataFrame(index=df.index)
    for col in compare_cols:
        keys[col] = df[col].map(lambda x: canonical_value(x, col))
    keep = ~keys.duplicated(keep="first")
    return df.loc[keep].reset_index(drop=True)


def find_col(columns, names):
    lookup = {norm_key(c): c for c in columns}
    for name in names:
        if norm_key(name) in lookup:
            return lookup[norm_key(name)]
    return None


def missing_mask(series, include_no=False):
    values = series.fillna("").astype(str).str.strip().str.casefold()
    allowed = set(MISSING_TEXT)
    if include_no:
        allowed |= {"no", "pending", "not available", "not nominated"}
    return values.isin(allowed)


def status_mask(df, col, pattern):
    if not col or col not in df.columns:
        return pd.Series(False, index=df.index)
    return df[col].fillna("").astype(str).str.strip().str.casefold().str.contains(pattern, regex=True, na=False)


def first_nonblank(series):
    for value in series:
        if pd.isna(value):
            continue
        text = str(value).strip()
        if text and text.casefold() not in MISSING_TEXT:
            return value
    return pd.NA


def employee_key(df, uin_col, name_col, bn_col):
    if uin_col and uin_col in df.columns:
        u = df[uin_col].fillna("").astype(str).str.strip()
    else:
        u = pd.Series("", index=df.index)
    name = df[name_col].fillna("").astype(str).str.strip().str.casefold() if name_col and name_col in df.columns else pd.Series("", index=df.index)
    bn = df[bn_col].fillna("").astype(str).str.strip().str.casefold() if bn_col and bn_col in df.columns else pd.Series("", index=df.index)
    fallback = "NAMEBN|" + name + "|" + bn
    key = u.where(u.ne(""), fallback)
    empty = key.eq("NAMEBN||") | key.eq("")
    key.loc[empty] = [f"ROW|{i}" for i in df.index[empty]]
    return key


def build_employee_master(df, uin_col, name_col, bn_col, status_col):
    if df.empty:
        return df.copy()
    work = df.copy()
    work["__employee_key__"] = employee_key(work, uin_col, name_col, bn_col)
    rows = []
    for _, group in work.groupby("__employee_key__", sort=False, dropna=False):
        row = {c: first_nonblank(group[c]) for c in df.columns}
        if status_col and status_col in group.columns:
            s = group[status_col].fillna("").astype(str).str.strip()
            nonblank = s[s.ne("")]
            if not nonblank.empty:
                completed = nonblank[nonblank.str.casefold().str.contains("complete|completed|done", regex=True, na=False)]
                row[status_col] = completed.iloc[0] if not completed.empty else nonblank.iloc[0]
        rows.append(row)
    return pd.DataFrame(rows, columns=df.columns).reset_index(drop=True)


def duplicate_uin_report(df, uin_col, name_col=None, bn_col=None):
    """Flag genuine UIN data-quality conflicts, not legitimate transaction history.
    A UIN is reported only when identity fields conflict across rows.
    Exact repeated records have already been removed.
    """
    if not uin_col or uin_col not in df.columns:
        return pd.DataFrame(columns=df.columns)
    work = df.copy()
    u = work[uin_col].fillna("").astype(str).str.strip()
    valid = work.loc[u.ne("")].copy()
    if valid.empty:
        return valid
    valid["__uin__"] = u.loc[valid.index]
    identity_cols = [c for c in [name_col, bn_col] if c and c in valid.columns]
    if not identity_cols:
        return pd.DataFrame(columns=df.columns)
    bad_uins = []
    for uid, g in valid.groupby("__uin__", sort=False):
        signatures = set()
        for _, r in g.iterrows():
            signatures.add(tuple(canonical_value(r[c], c) for c in identity_cols))
        if len(signatures) > 1:
            bad_uins.append(uid)
    if not bad_uins:
        return pd.DataFrame(columns=df.columns)
    return valid[valid["__uin__"].isin(bad_uins)].drop(columns=["__uin__"]).copy()


def completion_pct(done_or_missing, total, mode="missing"):
    if total <= 0:
        return 0.0
    value = ((total - done_or_missing) / total * 100) if mode == "missing" else (done_or_missing / total * 100)
    return round(max(0, min(100, value)), 1)


def money(v):
    try:
        return f"â‚¹{float(v):,.2f}"
    except Exception:
        return "â‚¹0.00"

# =========================================================
# EXCEL FORMATTING ENGINE
# =========================================================
def style_sheet(ws, header_row=1, start_row=2, end_row=None, start_col=1, end_col=None, add_table=False, table_name=None, freeze=None):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import ColorScaleRule

    if end_row is None:
        end_row = ws.max_row
    if end_col is None:
        end_col = ws.max_column
    if end_col < start_col:
        return

    navy = "17365D"
    blue = "1F4E78"
    light = "EAF2F8"
    border_color = "D6DEE8"
    green = "E2F0D9"
    yellow = "FFF2CC"

    thin = Side(style="thin", color=border_color)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(start_col, end_col + 1):
        cell = ws.cell(header_row, c)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[header_row].height = 30

    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name="Aptos", size=10, color="17202A")
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    for c in range(start_col, end_col + 1):
        letter = get_column_letter(c)
        max_len = 0
        for r in range(header_row, min(end_row, header_row + 250) + 1):
            v = ws.cell(r, c).value
            max_len = max(max_len, len(str(v)) if v is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 3, 11), 38)

    # Number/date formatting based on heading.
    for c in range(start_col, end_col + 1):
        heading = str(ws.cell(header_row, c).value or "").casefold()
        for r in range(start_row, end_row + 1):
            cell = ws.cell(r, c)
            if "date" in heading:
                cell.number_format = "dd-mm-yyyy"
            if any(x in heading for x in ["contribution", "amount", "total contribution"]):
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
            if "%" in heading or "percent" in heading:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.0%'

    if freeze:
        ws.freeze_panes = freeze
    else:
        ws.freeze_panes = f"A{header_row + 1}"

    ws.auto_filter.ref = f"{get_column_letter(start_col)}{header_row}:{get_column_letter(end_col)}{end_row}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_title_rows = f"${header_row}:${header_row}"

    if add_table and end_row >= start_row and end_col >= start_col:
        try:
            name = re.sub(r"[^A-Za-z0-9_]", "", table_name or f"Table{ws.title.replace(' ', '')}")
            if not name or name[0].isdigit():
                name = "TableReport"
            name = name[:200]
            tab = Table(displayName=name, ref=f"{get_column_letter(start_col)}{header_row}:{get_column_letter(end_col)}{end_row}")
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(tab)
        except Exception:
            pass


def make_simple_excel(df, sheet_name="Report", title=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export = df.copy()
        if title:
            title_df = pd.DataFrame([[title]], columns=["Report"])
            title_df.to_excel(writer, index=False, header=False, sheet_name=sheet_name, startrow=0)
            export.to_excel(writer, index=False, sheet_name=sheet_name, startrow=2)
            ws = writer.book[sheet_name]
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(export.columns)))
            ws["A1"].font = __import__("openpyxl").styles.Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
            ws["A1"].fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="17365D")
            ws["A1"].alignment = __import__("openpyxl").styles.Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 28
            style_sheet(ws, header_row=3, start_row=4, end_row=max(3, len(export) + 3), add_table=not export.empty, table_name="ReportTable")
        else:
            export.to_excel(writer, index=False, sheet_name=sheet_name)
            style_sheet(writer.book[sheet_name], add_table=not export.empty, table_name="ReportTable")
    output.seek(0)
    return output.getvalue()

# =========================================================
# UPLOAD
# =========================================================
with st.sidebar:
    st.markdown("### ðŸ“ Data Input")
    uploaded_files = st.file_uploader(
        "Upload Excel file(s)",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        help="Multiple Excel files and all worksheets are supported.",
    )
    st.caption("Up to 500 MB per file â€¢ XLSX / XLS")

if not uploaded_files:
    st.info("ðŸ‘† Upload one or more Excel files from the sidebar to start.")
    st.stop()

# =========================================================
# READ ALL FILES / SHEETS
# =========================================================
frames = []
source_log = []
errors = []

for file in uploaded_files:
    try:
        book = pd.ExcelFile(file)
        for sheet in book.sheet_names:
            try:
                raw = pd.read_excel(file, sheet_name=sheet)
                raw = clean_blank_rows(raw)
                if raw.empty:
                    continue
                raw.columns = normalize_columns(raw.columns)
                raw = merge_duplicate_columns(raw)
                raw["Source File"] = file.name
                raw["Source Sheet"] = sheet
                frames.append(raw)
                source_log.append({"Source File": file.name, "Source Sheet": sheet, "Records": len(raw)})
            except Exception as exc:
                errors.append(f"{file.name} / {sheet}: {exc}")
    except Exception as exc:
        errors.append(f"{file.name}: {exc}")

if errors:
    with st.expander("âš ï¸ File reading warnings"):
        for error in errors:
            st.warning(error)

if not frames:
    st.error("No readable records were found in the uploaded Excel files.")
    st.stop()

master_df = pd.concat(frames, ignore_index=True, sort=False)
master_df = clean_blank_rows(master_df)
raw_record_count = len(master_df)
master_df = deduplicate_exact_business_records(master_df)
removed_duplicate_records = raw_record_count - len(master_df)

# =========================================================
# COLUMN DETECTION
# =========================================================
uin_col = find_col(master_df.columns, ["UIN", "Employee Number", "Employee No", "Employee ID"])
name_col = find_col(master_df.columns, ["Employee Name", "Name", "Employee"])
bn_col = find_col(master_df.columns, ["BN", "Battalion", "Unit"])
status_col = find_col(master_df.columns, ["Case Status", "Status"])
pran_col = find_col(master_df.columns, ["PRAN No", "PRAN", "PRAN Number"])
nomination_col = find_col(master_df.columns, ["Nomination", "Nomination Status", "Nominee", "Nominee Status"])
amount_col = find_col(master_df.columns, ["Contribution", "Amount", "Total Amount"])

# =========================================================
# FILTERS
# =========================================================
st.sidebar.markdown("### ðŸ”Ž Filters")
search_text = st.sidebar.text_input("Search UIN / Name / BN / PRAN / Mobile / Status")

filtered_df = master_df.copy()
if search_text.strip():
    needle = re.escape(search_text.strip().casefold())
    text_df = filtered_df.fillna("").astype(str)
    mask = text_df.apply(lambda c: c.str.casefold().str.contains(needle, regex=True, na=False)).any(axis=1)
    filtered_df = filtered_df.loc[mask].copy()

if bn_col:
    bn_options = sorted(filtered_df[bn_col].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist())
    selected_bn = st.sidebar.selectbox("BN", ["All"] + bn_options)
    if selected_bn != "All":
        filtered_df = filtered_df[filtered_df[bn_col].astype(str).str.strip() == selected_bn].copy()
else:
    selected_bn = "All"

if status_col:
    status_options = sorted(filtered_df[status_col].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist())
    selected_status = st.sidebar.selectbox("Case Status", ["All"] + status_options)
    if selected_status != "All":
        filtered_df = filtered_df[filtered_df[status_col].astype(str).str.strip() == selected_status].copy()
else:
    selected_status = "All"

# =========================================================
# EMPLOYEE MASTER / QUALITY REPORTS
# =========================================================
employee_df = build_employee_master(filtered_df, uin_col, name_col, bn_col, status_col)
unique_employees = len(employee_df)

pending_count = int(status_mask(employee_df, status_col, r"pending").sum()) if status_col else 0
completed_count = int(status_mask(employee_df, status_col, r"complete|completed|done").sum()) if status_col else 0
missing_pran_df = employee_df[missing_mask(employee_df[pran_col])] if pran_col else pd.DataFrame(columns=employee_df.columns)
missing_nom_df = employee_df[missing_mask(employee_df[nomination_col], include_no=True)] if nomination_col else pd.DataFrame(columns=employee_df.columns)

duplicate_uin_df = duplicate_uin_report(filtered_df, uin_col, name_col, bn_col)

total_contribution = pd.to_numeric(filtered_df[amount_col], errors="coerce").fillna(0).sum() if amount_col else 0
pran_completion = completion_pct(len(missing_pran_df), unique_employees)
nomination_completion = completion_pct(len(missing_nom_df), unique_employees)

# =========================================================
# TOP STATUS BAR
# =========================================================
st.success(
    f"Loaded {len(uploaded_files)} file(s) â€¢ {len(source_log)} worksheet(s) â€¢ "
    f"{len(master_df):,} unique business records"
)
if removed_duplicate_records:
    st.info(f"â™»ï¸ {removed_duplicate_records:,} exact duplicate record(s) were removed automatically. Source File/Source Sheet were ignored during duplicate comparison.")

with st.expander("ðŸ“ Uploaded Files Summary"):
    st.dataframe(pd.DataFrame(source_log), use_container_width=True, hide_index=True)

# =========================================================
# DASHBOARD
# =========================================================
st.markdown("## ðŸ“ˆ Automatic Dashboard")
st.caption("Employee-level KPIs use one consolidated row per employee. Contribution uses the cleaned transaction/history records.")

r1 = st.columns(5)
r1[0].metric("Total Records", f"{len(filtered_df):,}")
r1[1].metric("Unique Employees", f"{unique_employees:,}")
r1[2].metric("Pending", f"{pending_count:,}")
r1[3].metric("Completed", f"{completed_count:,}")
r1[4].metric("Duplicate UIN", f"{len(duplicate_uin_df):,}")

r2 = st.columns(5)
r2[0].metric("Missing PRAN", f"{len(missing_pran_df):,}")
r2[1].metric("PRAN Completion", f"{pran_completion:.1f}%")
r2[2].metric("Missing Nomination", f"{len(missing_nom_df):,}")
r2[3].metric("Nomination Completion", f"{nomination_completion:.1f}%")
r2[4].metric("Total Contribution", money(total_contribution))

# =========================================================
# BN SUMMARY
# =========================================================
bn_summary = pd.DataFrame()
if bn_col:
    st.markdown("## ðŸ¢ BN-wise Summary")
    rows = []
    for bn_value, emp_group in employee_df.groupby(bn_col, dropna=False, sort=True):
        display_bn = str(bn_value).strip() if not pd.isna(bn_value) else ""
        display_bn = display_bn if display_bn and display_bn.casefold() not in {"nan", "none"} else "Blank BN"
        transaction_group = filtered_df[filtered_df[bn_col].fillna("").astype(str).str.strip() == ("" if pd.isna(bn_value) else str(bn_value).strip())]
        contribution = pd.to_numeric(transaction_group[amount_col], errors="coerce").fillna(0).sum() if amount_col else 0
        rows.append({
            "BN": display_bn,
            "Total Employees": len(emp_group),
            "Pending": int(status_mask(emp_group, status_col, r"pending").sum()) if status_col else 0,
            "Completed": int(status_mask(emp_group, status_col, r"complete|completed|done").sum()) if status_col else 0,
            "Missing PRAN": int(missing_mask(emp_group[pran_col]).sum()) if pran_col else 0,
            "Missing Nomination": int(missing_mask(emp_group[nomination_col], include_no=True).sum()) if nomination_col else 0,
            "Duplicate UIN Records": len(duplicate_uin_report(transaction_group, uin_col, name_col, bn_col)),
            "Total Contribution": contribution,
        })
    bn_summary = pd.DataFrame(rows)
    if not bn_summary.empty:
        for metric in ["Pending", "Completed", "Missing PRAN", "Missing Nomination"]:
            bn_summary[f"{metric} %"] = (bn_summary[metric] / bn_summary["Total Employees"].replace(0, pd.NA) * 100).fillna(0).round(1)
        st.dataframe(bn_summary, use_container_width=True, hide_index=True)
        chart_cols = [c for c in ["Total Employees", "Pending", "Completed", "Missing PRAN", "Missing Nomination"] if c in bn_summary.columns]
        if chart_cols:
            st.bar_chart(bn_summary.set_index("BN")[chart_cols], use_container_width=True)

# =========================================================
# CASE STATUS
# =========================================================
if status_col:
    st.markdown("## ðŸ“Œ Case Status Summary")
    status_summary = filtered_df[status_col].fillna("Blank").astype(str).str.strip().replace("", "Blank").value_counts().rename_axis("Case Status").reset_index(name="Total")
    st.dataframe(status_summary, use_container_width=True, hide_index=True)
else:
    status_summary = pd.DataFrame()

# =========================================================
# INDIVIDUAL REPORTS
# =========================================================
st.markdown("## ðŸ‘¤ Employee-wise Report")
st.write(f"**Total Unique Employees: {unique_employees:,}**")
st.dataframe(employee_df, use_container_width=True, hide_index=True, height=420)
st.download_button("â¬‡ï¸ Download Employee-wise Excel", make_simple_excel(employee_df, "Employee Report", "Employee-wise Master Report"), "Employee_Wise_Report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="employee_download")

st.markdown("## âš ï¸ Duplicate UIN Report")
if duplicate_uin_df.empty:
    st.success("âœ… No conflicting duplicate UIN found.")
else:
    st.warning(f"{len(duplicate_uin_df):,} record(s) require UIN verification.")
    st.dataframe(duplicate_uin_df, use_container_width=True, hide_index=True)
    st.download_button("â¬‡ï¸ Download Duplicate UIN Report", make_simple_excel(duplicate_uin_df, "Duplicate UIN", "Duplicate / Conflicting UIN Report"), "Duplicate_UIN_Report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dup_download")

st.markdown("## ðŸªª Missing PRAN Report")
if missing_pran_df.empty:
    st.success("âœ… All employees have PRAN.")
else:
    st.warning(f"Employees without PRAN: {len(missing_pran_df):,}")
    st.dataframe(missing_pran_df, use_container_width=True, hide_index=True)
    st.download_button("â¬‡ï¸ Download Missing PRAN Report", make_simple_excel(missing_pran_df, "Missing PRAN", "Employees Without PRAN"), "Missing_PRAN_Report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="pran_download")

st.markdown("## ðŸ“ Missing Nomination Report")
if missing_nom_df.empty:
    st.success("âœ… All employees have nomination.")
else:
    st.warning(f"Employees with missing nomination: {len(missing_nom_df):,}")
    st.dataframe(missing_nom_df, use_container_width=True, hide_index=True)
    st.download_button("â¬‡ï¸ Download Missing Nomination Report", make_simple_excel(missing_nom_df, "Missing Nomination", "Employees With Missing Nomination"), "Missing_Nomination_Report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="nom_download")

# =========================================================
# EXCEL REPORT BUILDERS
# =========================================================
def write_title(ws, title, end_col):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, end_col))
    cell = ws.cell(1, 1)
    cell.value = title
    cell.font = Font(name="Aptos Display", size=17, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32


def write_summary_sheet(writer, summary_df, title="Smart Excel Reporting Agent"):
    summary_df.to_excel(writer, index=False, sheet_name="Summary", startrow=2)
    ws = writer.book["Summary"]
    write_title(ws, title, 2)
    style_sheet(ws, header_row=3, start_row=4, end_row=len(summary_df) + 3, end_col=2, add_table=not summary_df.empty, table_name="SummaryTable")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 24


def write_table_sheet(writer, df, sheet_name, title=None):
    safe = re.sub(r"[\[\]:*?/\\]", "_", sheet_name)[:31] or "Report"
    df.to_excel(writer, index=False, sheet_name=safe, startrow=2 if title else 0)
    ws = writer.book[safe]
    if title:
        write_title(ws, title, max(2, len(df.columns)))
        header_row = 3
        start_row = 4
    else:
        header_row = 1
        start_row = 2
    end_row = header_row + len(df)
    style_sheet(ws, header_row=header_row, start_row=start_row, end_row=max(header_row, end_row), add_table=not df.empty, table_name=f"Tbl{re.sub(r'[^A-Za-z0-9]', '', safe)}")
    return safe


def build_complete_report(include_bn_sheets=False):
    output = BytesIO()
    summary = pd.DataFrame({
        "Particular": [
            "Total Records", "Unique Employees", "Pending", "Completed",
            "Missing PRAN", "PRAN Completion", "Missing Nomination",
            "Nomination Completion", "Duplicate UIN Records", "Total Contribution",
            "Uploaded Files", "Duplicate Records Removed"
        ],
        "Value": [
            len(filtered_df), unique_employees, pending_count, completed_count,
            len(missing_pran_df), f"{pran_completion:.1f}%", len(missing_nom_df),
            f"{nomination_completion:.1f}%", len(duplicate_uin_df), total_contribution,
            len(uploaded_files), removed_duplicate_records
        ]
    })
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        write_summary_sheet(writer, summary)
        if not bn_summary.empty:
            write_table_sheet(writer, bn_summary, "BN Summary", "BN-wise Summary")
        if not status_summary.empty:
            write_table_sheet(writer, status_summary, "Case Status", "Case Status Summary")
        write_table_sheet(writer, employee_df, "Employee Report", "Employee-wise Master Report")
        write_table_sheet(writer, missing_pran_df, "Missing PRAN", "Employees Without PRAN")
        write_table_sheet(writer, missing_nom_df, "Missing Nomination", "Employees With Missing Nomination")
        write_table_sheet(writer, duplicate_uin_df, "Duplicate UIN", "Duplicate / Conflicting UIN")
        write_table_sheet(writer, filtered_df, "Filtered Data", "Cleaned Filtered Data")
        if include_bn_sheets and bn_col:
            used = {"Summary", "BN Summary", "Case Status", "Employee Report", "Missing PRAN", "Missing Nomination", "Duplicate UIN", "Filtered Data"}
            for bn_value, group in filtered_df.groupby(bn_col, dropna=False, sort=True):
                bn_name = str(bn_value).strip() if not pd.isna(bn_value) else "Blank BN"
                safe = re.sub(r"[\[\]:*?/\\]", "_", f"BN-{bn_name}")[:31] or "BN"
                base = safe
                n = 1
                while safe in used:
                    suffix = f"_{n}"
                    safe = base[:31-len(suffix)] + suffix
                    n += 1
                used.add(safe)
                # Summary block
                mini = bn_summary[bn_summary["BN"].astype(str) == bn_name].copy()
                if mini.empty:
                    mini = pd.DataFrame({"Particular": ["BN"], "Value": [bn_name]})
                mini2 = pd.DataFrame({"Particular": ["BN", "Total Employees", "Pending", "Completed", "Missing PRAN", "Missing Nomination", "Duplicate UIN Records", "Total Contribution"], "Value": [bn_name, len(build_employee_master(group, uin_col, name_col, bn_col, status_col)), int(status_mask(build_employee_master(group, uin_col, name_col, bn_col, status_col), status_col, r"pending").sum()) if status_col else 0, int(status_mask(build_employee_master(group, uin_col, name_col, bn_col, status_col), status_col, r"complete|completed|done").sum()) if status_col else 0, int(missing_mask(build_employee_master(group, uin_col, name_col, bn_col, status_col)[pran_col]).sum()) if pran_col else 0, int(missing_mask(build_employee_master(group, uin_col, name_col, bn_col, status_col)[nomination_col], include_no=True).sum()) if nomination_col else 0, len(duplicate_uin_report(group, uin_col, name_col, bn_col)), pd.to_numeric(group[amount_col], errors="coerce").fillna(0).sum() if amount_col else 0]})
                mini2.to_excel(writer, index=False, sheet_name=safe, startrow=2)
                group.to_excel(writer, index=False, sheet_name=safe, startrow=len(mini2) + 4)
                ws = writer.book[safe]
                write_title(ws, f"BN-wise Detailed Report - {bn_name}", len(group.columns))
                style_sheet(ws, header_row=3, start_row=4, end_row=3+len(mini2), end_col=2, add_table=True, table_name=f"Summary{len(used)}")
                data_header = len(mini2) + 5
                style_sheet(ws, header_row=data_header, start_row=data_header+1, end_row=data_header+len(group), end_col=max(1, len(group.columns)), add_table=not group.empty, table_name=f"Data{len(used)}", freeze=f"A{data_header+1}")
    output.seek(0)
    return output.getvalue()

# =========================================================
# DOWNLOAD CENTRE
# =========================================================
st.markdown("## ðŸ“¦ Excel Report Centre")
st.caption("All reports are formatted with professional headers, borders, filters, freeze panes, automatic widths, print settings and Excel table styling.")

c1, c2 = st.columns(2)
with c1:
    st.download_button(
        "â¬‡ï¸ Download Complete Master Report",
        build_complete_report(False),
        "Complete_Master_Report.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="master_report_download",
        use_container_width=True,
    )
with c2:
    st.download_button(
        "â¬‡ï¸ Download BN-wise Complete Report",
        build_complete_report(True),
        "BN_Wise_Complete_Report.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="bn_report_download",
        use_container_width=True,
    )

st.download_button(
    "â¬‡ï¸ Download Cleaned Filtered Excel",
    make_simple_excel(filtered_df, "Filtered Data", "Cleaned Filtered Data"),
    "Filtered_Report.xlsx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    key="filtered_download",
    use_container_width=True,
)

# =========================================================
# PREVIEW
# =========================================================
st.markdown("## ðŸ‘€ Data Preview")
st.dataframe(filtered_df, use_container_width=True, height=450, hide_index=True)

with st.expander("â„¹ï¸ Detected Columns"):
    st.write({
        "UIN": uin_col,
        "Employee Name": name_col,
        "BN": bn_col,
        "PRAN": pran_col,
        "Nomination": nomination_col,
        "Case Status": status_col,
        "Contribution": amount_col,
    })
